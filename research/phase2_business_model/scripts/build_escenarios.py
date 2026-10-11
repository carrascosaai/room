#!/usr/bin/env python3
"""Genera escenarios.xlsx (fórmulas editables) y escenarios_resultados.csv (mismos cálculos en Python).

Uso: python3 research/phase2_business_model/scripts/build_escenarios.py
Edita SUPUESTOS aquí o directamente la hoja 'Supuestos' del xlsx (las demás hojas son fórmulas).
Ningún escenario es una previsión: los supuestos marcados "supuesto" no tienen fuente.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).resolve().parent.parent
SC = ["conservador", "base", "optimista"]

# clave, descripción, unidad, (conservador, base, optimista), tipo, fuente
SUPUESTOS = [
    ("act", "Usuarios activos sobre registrados", "fracción", (0.30, 0.40, 0.60), "supuesto", "—"),
    ("conv", "Conversión de activos a pago (A/F)", "fracción", (0.01, 0.03, 0.05), "supuesto", "sin datos de competidores; medir con phase2_customers/07"),
    ("precio_a", "Precio suscripción estudiante (A)", "€/mes", (3, 5, 9), "hipótesis a probar", "phase2_customers/07 (0/3/5/9 €)"),
    ("meses_a", "Meses pagados por año (estacionalidad)", "meses", (2, 3, 5), "supuesto", "Simplify vende semanal/trimestral (fase 1, [T])"),
    ("precio_f", "Precio tracker (F)", "€/mes", (3, 4, 6), "supuesto", "JobFill 3,99 $; GoFillr 5,99 $; Huntr 40 $ (fase 1)"),
    ("conv_f", "Conversión a tracker de pago (F)", "fracción", (0.005, 0.01, 0.02), "supuesto", "—"),
    ("precio_oferta", "Precio publicación por oferta (B)", "€/oferta", (20, 30, 49), "dato de mercado", "EUInternships 20–39 €; iAgora desde 23 €; InfoJobs 49 € 1.ª oferta [T]"),
    ("ofertas_cliente", "Ofertas publicadas por empresa y mes (B)", "nº/mes", (1, 1.5, 3), "supuesto", "—"),
    ("fee_colocacion", "Honorario por incorporación en prácticas o junior (C)", "€/incorporación", (300, 800, 1500), "supuesto acotado", "headhunting 15–33 % salario [T]; junior ≈26–32 k€ [T]; Talentoteca 12 % de la beca"),
    ("coloc_cliente_ano", "Incorporaciones por empresa cliente y año (C)", "nº/año", (1, 2, 3), "supuesto", "—"),
    ("precio_destacada", "Oferta patrocinada/destacada (D)", "€/oferta/mes", (30, 60, 100), "supuesto acotado", "LinkedIn ~10 $/día mínimo; StepStone Pro 1.449 €/30 días [T]"),
    ("destacadas_cliente", "Destacadas por empresa y mes (D)", "nº/mes", (1, 1, 2), "supuesto", "—"),
    ("licencia_uni", "Licencia anual universidad o escuela (E)", "€/año", (500, 1500, 4000), "supuesto", "JobTeaser núcleo gratis para centros [C]; Symplicity £12.050/licencia (G-Cloud) [T]"),
    ("herr_fijas", "Herramientas fijas (dominio, hosting, BD, email, analítica)", "€/mes", (5, 30, 80), "estimación", "Supabase Pro 25 $; Resend Pro 20 $; dominio ~1,5 €/mes [T]"),
    ("gestoria", "Gestoría", "€/mes", (0, 50, 80), "supuesto", "—"),
    ("cuota_autonomo", "Cuota de autónomos", "€/mes", (88.64, 88.64, 250), "dato secundario", "tarifa plana 80 € + MEI ≈ 88,64 € el 1.er año; tramos ~206–623 € [T]"),
    ("publicidad", "Publicidad", "€/mes", (0, 100, 400), "decisión", "phase2_growth"),
    ("comision_pago", "Comisión de pasarela sobre ingresos", "fracción", (0.03, 0.03, 0.03), "supuesto", "orden de magnitud de pasarelas con tarjeta; verificar"),
    ("horas_mant", "Horas del fundador al mes en mantener datos (30 empresas)", "h/mes", (59, 42, 27), "estimación", "phase2_technical/data/coste_resultados.csv"),
    ("horas_comercial", "Horas del fundador al mes en ventas y soporte", "h/mes", (40, 40, 40), "supuesto", "—"),
    ("irpf", "Tipo efectivo IRPF sobre el beneficio (autónomo)", "fracción", (0.10, 0.15, 0.20), "supuesto simplificado", "no es asesoramiento fiscal"),
    ("clientes_b", "Empresas cliente (B) en el escenario de PyG", "nº", (3, 10, 20), "supuesto", "—"),
    ("clientes_d", "Empresas cliente (D) en el escenario de PyG", "nº", (2, 6, 10), "supuesto", "—"),
    ("clientes_c", "Empresas cliente (C) en el escenario de PyG", "nº", (1, 4, 6), "supuesto", "—"),
    ("unis_e", "Universidades o escuelas cliente (E) en PyG", "nº", (0, 1, 2), "supuesto", "—"),
    ("usuarios_pyg", "Usuarios registrados en el escenario de PyG", "nº", (300, 1500, 6000), "supuesto", "—"),
]
USERS = [100, 1000, 10000]
TARGETS = [1000, 5000, 10000]


def compute(s):
    """Mismos cálculos que las fórmulas del xlsx, para el CSV de resultados."""
    out = {}
    for u in USERS:
        out[f"A_MRR_temporada_{u}"] = u * s["act"] * s["conv"] * s["precio_a"]
        out[f"A_anual_{u}"] = out[f"A_MRR_temporada_{u}"] * s["meses_a"]
        out[f"F_anual_{u}"] = u * s["act"] * s["conv_f"] * s["precio_f"] * 12
    arpa = {"B": s["precio_oferta"] * s["ofertas_cliente"], "C": s["fee_colocacion"] * s["coloc_cliente_ano"] / 12,
            "D": s["precio_destacada"] * s["destacadas_cliente"], "E": s["licencia_uni"] / 12}
    for m, v in arpa.items():
        out[f"ARPA_mensual_{m}"] = v
        for t in TARGETS:
            out[f"clientes_{m}_para_{t}"] = -(-t // v)  # techo
    ingresos = (12 * (arpa["B"] * s["clientes_b"] + arpa["D"] * s["clientes_d"] + arpa["C"] * s["clientes_c"] + arpa["E"] * s["unis_e"])
                + s["usuarios_pyg"] * s["act"] * s["conv"] * s["precio_a"] * s["meses_a"])
    gastos = 12 * (s["herr_fijas"] + s["gestoria"] + s["cuota_autonomo"] + s["publicidad"]) + ingresos * s["comision_pago"]
    bai = ingresos - gastos
    neto = bai - max(bai, 0) * s["irpf"]
    horas = 12 * (s["horas_mant"] + s["horas_comercial"])
    out.update({"PyG_facturacion_anual": ingresos, "PyG_gastos_anuales": gastos, "PyG_BAI": bai, "PyG_IRPF": max(bai, 0) * s["irpf"],
                "PyG_beneficio_neto": neto, "PyG_sueldo_fundador_anual": max(neto, 0), "PyG_euros_por_hora_fundador": neto / horas})
    return out


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Supuestos"
    bold = Font(bold=True)
    edit = PatternFill("solid", fgColor="FFF2CC")
    ws.append(["clave", "descripción", "unidad", *SC, "tipo", "fuente"])
    for c in ws[1]:
        c.font = bold
    ref = {}
    for i, (k, d, u, vals, t, f) in enumerate(SUPUESTOS, start=2):
        ws.append([k, d, u, *vals, t, f])
        for col in "DEF":
            ws[f"{col}{i}"].fill = edit
        ref[k] = i
    ws.append([])
    ws.append(["Las celdas amarillas son editables. Las demás hojas se recalculan con fórmulas."])
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["H"].width = 60

    def R(k, col):
        return f"Supuestos!${col}${ref[k]}"

    cols = {"conservador": "D", "base": "E", "optimista": "F"}

    # Hoja B2C
    b = wb.create_sheet("B2C_usuarios")
    b.append(["Ingresos de estudiantes (A: suscripción; F: tracker) según usuarios registrados"])
    b.append(["usuarios", "escenario", "MRR en temporada A (€)", "Ingresos anuales A (€)", "Ingresos anuales F (€)"])
    for u in USERS:
        for sc, col in cols.items():
            r = b.max_row + 1
            b.append([u, sc,
                      f"=A{r}*{R('act', col)}*{R('conv', col)}*{R('precio_a', col)}",
                      f"=C{r}*{R('meses_a', col)}",
                      f"=A{r}*{R('act', col)}*{R('conv_f', col)}*{R('precio_f', col)}*12"])

    # Hoja B2B
    e = wb.create_sheet("B2B_clientes")
    e.append(["Clientes empresariales necesarios para cada facturación mensual objetivo"])
    e.append(["modelo", "escenario", "ARPA mensual (€)", *[f"clientes para {t} €/mes" for t in TARGETS]])
    arpa_f = {"B (publicación)": lambda c: f"{R('precio_oferta', c)}*{R('ofertas_cliente', c)}",
              "C (captación/colocación)": lambda c: f"{R('fee_colocacion', c)}*{R('coloc_cliente_ano', c)}/12",
              "D (patrocinadas)": lambda c: f"{R('precio_destacada', c)}*{R('destacadas_cliente', c)}",
              "E (universidad)": lambda c: f"{R('licencia_uni', c)}/12"}
    for m, fn in arpa_f.items():
        for sc, col in cols.items():
            r = e.max_row + 1
            e.append([m, sc, f"={fn(col)}", *[f"=ROUNDUP({t}/C{r},0)" for t in TARGETS]])

    # Hoja PyG
    p = wb.create_sheet("PyG")
    p.append(["Cuenta de resultados anual (autónomo persona física; simplificada, no es asesoramiento fiscal)"])
    p.append(["concepto", *SC])
    lines = [
        ("Facturación B (publicación)", lambda c: f"12*{R('precio_oferta', c)}*{R('ofertas_cliente', c)}*{R('clientes_b', c)}"),
        ("Facturación D (patrocinadas)", lambda c: f"12*{R('precio_destacada', c)}*{R('destacadas_cliente', c)}*{R('clientes_d', c)}"),
        ("Facturación C (colocación)", lambda c: f"{R('fee_colocacion', c)}*{R('coloc_cliente_ano', c)}*{R('clientes_c', c)}"),
        ("Facturación E (universidades)", lambda c: f"{R('licencia_uni', c)}*{R('unis_e', c)}"),
        ("Facturación A (estudiantes)", lambda c: f"{R('usuarios_pyg', c)}*{R('act', c)}*{R('conv', c)}*{R('precio_a', c)}*{R('meses_a', c)}"),
    ]
    for name, fn in lines:
        p.append([name, *[f"={fn(c)}" for c in cols.values()]])
    p.append(["FACTURACIÓN TOTAL", *[f"=SUM({c}3:{c}7)" for c in "BCD"]])
    p.append(["Herramientas", *[f"=12*{R('herr_fijas', c)}" for c in cols.values()]])
    p.append(["Gestoría", *[f"=12*{R('gestoria', c)}" for c in cols.values()]])
    p.append(["Cuota de autónomos", *[f"=12*{R('cuota_autonomo', c)}" for c in cols.values()]])
    p.append(["Publicidad", *[f"=12*{R('publicidad', c)}" for c in cols.values()]])
    p.append(["Comisiones de pago", *[f"={cc}8*{R('comision_pago', c)}" for cc, c in zip("BCD", cols.values())]])
    p.append(["GASTOS TOTALES", *[f"=SUM({c}9:{c}13)" for c in "BCD"]])
    p.append(["BENEFICIO ANTES DE IMPUESTOS", *[f"={c}8-{c}14" for c in "BCD"]])
    p.append(["IRPF estimado", *[f"=MAX({cc}15,0)*{R('irpf', c)}" for cc, c in zip("BCD", cols.values())]])
    p.append(["BENEFICIO NETO", *[f"={c}15-{c}16" for c in "BCD"]])
    p.append(["SUELDO DEL FUNDADOR (autónomo = beneficio neto, si es >0)", *[f"=MAX({c}17,0)" for c in "BCD"]])
    p.append(["Horas del fundador al año", *[f"=12*({R('horas_mant', c)}+{R('horas_comercial', c)})" for c in cols.values()]])
    p.append(["€/hora implícito del fundador", *[f"={c}17/{c}19" for c in "BCD"]])
    for row in p.iter_rows(min_row=2, max_row=2):
        for cell in row:
            cell.font = bold
    p.column_dimensions["A"].width = 58
    wb.save(OUT / "escenarios.xlsx")

    # CSV con los mismos cálculos
    with open(OUT / "escenarios_resultados.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        res = {sc: compute({k: v[i] for k, _, _, v, _, _ in SUPUESTOS}) for i, sc in enumerate(SC)}
        w.writerow(["metrica", *SC])
        for k in res["base"]:
            w.writerow([k, *[round(res[sc][k], 2) for sc in SC]])
    print("escenarios.xlsx y escenarios_resultados.csv generados")


if __name__ == "__main__":
    main()
